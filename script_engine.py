"""
KoKoFish — Script Lab engine.

Handles character profile management, script parsing,
AI script tagging, and multi-voice playback coordination.
"""

import json
import logging
import os
import re
import sys

logger = logging.getLogger("KoKoFish.script")

APP_DIR = (
    os.path.dirname(sys.executable)
    if getattr(sys, "frozen", False)
    else os.path.dirname(os.path.abspath(__file__))
)
PROFILES_DIR = os.path.join(APP_DIR, "scripts", "profiles")
SCRIPTS_DIR  = os.path.join(APP_DIR, "scripts")

# Engine name → subfolder mapping so profiles stay separated
_ENGINE_FOLDER = {
    "kokoro":     "kokoro",
    "voxcpm_05b": "voxcpm_05b",
    "voxcpm_2b":  "voxcpm_2b",
    "omnivoice":  "omnivoice",
}


# ---------------------------------------------------------------------------
# Profile management
# ---------------------------------------------------------------------------

def get_profiles_dir(engine: str = "kokoro") -> str:
    folder = _ENGINE_FOLDER.get(engine, engine) or "kokoro"
    d = os.path.join(PROFILES_DIR, folder)
    os.makedirs(d, exist_ok=True)
    return d


def list_profiles(engine: str = "kokoro") -> list:
    d = get_profiles_dir(engine)
    return [f[:-5] for f in sorted(os.listdir(d)) if f.endswith(".json")]


def load_profile(name: str, engine: str = "kokoro") -> dict:
    path = os.path.join(get_profiles_dir(engine), f"{name}.json")
    if os.path.isfile(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as exc:
            logger.warning("load_profile failed: %s", exc)
    return default_profile()


def save_profile(name: str, profile: dict, engine: str = "kokoro"):
    d = get_profiles_dir(engine)
    path = os.path.join(d, f"{name}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(profile, f, indent=2, ensure_ascii=False)
    logger.info("Profile saved: %s (engine=%s)", name, engine)


def delete_profile(name: str, engine: str = "kokoro"):
    path = os.path.join(get_profiles_dir(engine), f"{name}.json")
    if os.path.isfile(path):
        os.remove(path)
        logger.info("Profile deleted: %s (engine=%s)", name, engine)


def default_profile() -> dict:
    return {
        "characters": [],
        "narrator": {"voice": "", "blend_voice": "", "blend_ratio": 0.0},
    }


# ---------------------------------------------------------------------------
# Script parsing
# ---------------------------------------------------------------------------

def parse_script(text: str) -> list:
    """
    Parse a [Character] tagged script into a list of dicts.
    Each dict: {character: str, text: str}

    Accepts both:
      [Name] dialogue text
      [Name]: "dialogue text"   (colon + optional surrounding quotes)

    Lines without a tag are assigned to Narrator.
    """
    segments = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.match(r'^\[([^\]]+)\]\s*:?\s*(.*)', line)
        if m:
            char = m.group(1).strip()
            txt  = m.group(2).strip()
            # Strip surrounding straight or curly quotes
            if len(txt) >= 2 and txt[0] in ('"', '\u201c') and txt[-1] in ('"', '\u201d'):
                txt = txt[1:-1]
            segments.append({"character": char, "text": txt})
        else:
            segments.append({"character": "Narrator", "text": line})
    return segments


def format_script(segments: list) -> str:
    """Convert parsed segments back to tagged script text."""
    return "\n".join(f"[{s['character']}] {s['text']}" for s in segments)


def read_source_file(path: str) -> str:
    """Read a source document (txt, pdf, docx, epub) to plain text."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".txt":
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        elif ext == ".pdf":
            import pdfplumber
            pages = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        pages.append(t)
            return "\n\n".join(pages)
        elif ext == ".docx":
            from docx import Document
            doc = Document(path)
            return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext == ".epub":
            import ebooklib
            from ebooklib import epub
            from html.parser import HTMLParser

            class _Strip(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.parts = []
                def handle_data(self, data):
                    self.parts.append(data)

            book = epub.read_epub(path)
            chapters = []
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                p = _Strip()
                p.feed(item.get_content().decode("utf-8", errors="replace"))
                chapters.append("".join(p.parts))
            return "\n\n".join(chapters)
    except Exception as exc:
        raise RuntimeError(f"Could not read {os.path.basename(path)}: {exc}")
    return ""


# ---------------------------------------------------------------------------
# AI script tagging prompts
# ---------------------------------------------------------------------------

_KOKORO_PROMPT = """\
You are converting prose fiction into a voice-acted script for Kokoro TTS (plain text only).

OUTPUT FORMAT — every line must be exactly:
[CharacterName] the words they say or the narration

RULES:
1. SPOKEN DIALOGUE (text inside quotation marks):
   - Identify the speaker from nearby attribution: "Maren said", "she replied", "Vehn asked"
   - Pronouns (she/he/they) refer to the most recently named character of that gender in the text
   - Write: [SpeakerName] the exact words spoken — no quotation marks, no attribution phrase
   - REMOVE attribution entirely: "said", "replied", "asked", "whispered", "answered", "commented", "continued", "interjected", "added", "chirped", "nodded" etc.

2. NARRATION (prose that is not inside quotation marks):
   - Write: [Narrator] the narration exactly as written

3. EVERY line must start with a [Name] tag — no exceptions, no blank lines without tags
4. Use the EXACT words from the source — never summarize, paraphrase, or invent content
5. A name inside quotation marks ("She called to John") is a reference, NOT a speaker — ignore it for attribution
6. When the speaker is genuinely unclear → [Narrator]
7. Do NOT add emotion tags, brackets, ellipses, or any formatting not in the source
8. Known characters: {characters}

EXAMPLE:
Input:
  "Ready when you are," Maren said quietly. Vehn nodded toward the console.
  "Understood," he replied.
Output:
  [Maren] Ready when you are.
  [Narrator] Vehn nodded toward the console.
  [Vehn] Understood.

Output ONLY the formatted script lines. No explanations, no commentary, no placeholders.\
"""

_FISH_PROMPT = """\
You are converting prose fiction into a voice-acted script for Fish-Speech TTS.

OUTPUT FORMAT — every line must be exactly:
[CharacterName] (emotion) the words they say or the narration

RULES:
1. SPOKEN DIALOGUE (text inside quotation marks):
   - Identify the speaker from nearby attribution: "Maren said", "she replied", "Vehn asked"
   - Pronouns (she/he/they) refer to the most recently named character of that gender in the text
   - Write: [SpeakerName] (emotion) the exact words spoken — no quotation marks, no attribution phrase
   - REMOVE attribution entirely: "said", "replied", "asked", "whispered", "answered", "commented", "continued", "interjected", "added" etc.
   - Choose ONE emotion tag: (excited) (happy) (satisfied) (confident) (gentle) (serious) (sad) (angry) (nervous) (fearful) (surprised) (confused) (laugh) [whisper]

2. NARRATION (prose that is not inside quotation marks):
   - Write: [Narrator] the narration exactly as written — no emotion tag on narration

3. EVERY line must start with a [Name] tag — no exceptions
4. Use the EXACT words from the source — never summarize, paraphrase, or invent content
5. A name inside quotation marks ("She called to John") is a reference, NOT a speaker — ignore it for attribution
6. When the speaker is genuinely unclear → [Narrator]
7. Match emotions to conversational flow — consider what was just said before choosing the emotion
8. Known characters: {characters}

EXAMPLE:
Input:
  "Ready when you are," Maren said quietly. Vehn nodded toward the console.
  "Understood," he replied.
Output:
  [Maren] (gentle) Ready when you are.
  [Narrator] Vehn nodded toward the console.
  [Vehn] (confident) Understood.

Output ONLY the formatted script lines. No explanations, no commentary, no placeholders.\
"""

_ENHANCE_KOKORO_PROMPT = (
    "You are a script editor improving a voice-acted script for natural flow and emotional continuity.\n"
    "The script uses the format: [CharacterName] dialogue text\n"
    "Rules:\n"
    "1. Keep all [CharacterName] tags exactly as they are. Do NOT change speaker names.\n"
    "2. Improve the dialogue text for natural spoken delivery — fix awkward phrasing, "
    "remove written-prose habits, and make lines sound like they were actually spoken.\n"
    "3. Ensure emotional continuity: each line should feel like a natural response to the previous one. "
    "Characters should react to each other realistically.\n"
    "4. Do NOT add emotion tags — this is for the Kokoro engine which uses plain text.\n"
    "5. Do NOT summarize, shorten, or remove lines. Every line in must produce a line out.\n"
    "6. IMPORTANT: Names inside quotation marks within dialogue are references, not speakers — leave them.\n"
    "Output ONLY the improved script. No explanations, no commentary."
)

_ENHANCE_FISH_PROMPT = (
    "You are a script editor improving a voice-acted script for natural flow and emotional continuity.\n"
    "The script uses the format: [CharacterName] (emotion) dialogue text\n"
    "Rules:\n"
    "1. Keep all [CharacterName] tags exactly as they are. Do NOT change speaker names.\n"
    "2. Update or correct the (emotion) tag on each dialogue line so it matches the flow of conversation. "
    "Valid tags: (excited) (happy) (satisfied) (confident) (gentle) (serious) (sad) (angry) "
    "(nervous) (fearful) (surprised) (confused) (laugh) [whisper]\n"
    "3. Improve dialogue text for natural spoken delivery — fix awkward phrasing and written-prose habits.\n"
    "4. Ensure emotional continuity: each character's emotional state should react realistically to the previous line.\n"
    "5. Do NOT summarize, shorten, or remove lines. Every line in must produce a line out.\n"
    "6. IMPORTANT: Names inside quotation marks within dialogue are references, not speakers — leave them.\n"
    "Output ONLY the improved script. No explanations, no commentary."
)


_SKIP_WORDS = frozenset({
    "The", "She", "He", "They", "It", "This", "There", "But", "And", "So",
    "Then", "When", "That", "What", "Which", "Who", "Where", "How", "Now",
    "Just", "Still", "While", "After", "Before", "With", "Into", "Onto",
    "Alright", "Already", "Always", "Also", "Again", "Once",
})

_ATTR_VERBS = re.compile(
    r'\b([A-Z][a-z]{2,})\s+'
    r'(?:said|says|replied|replies|asked|asks|whispered|whispers|'
    r'shouted|shouts|called|calls|answered|answers|commented|comments|'
    r'continued|continues|interjects|interjected|chirped|chirps|'
    r'nodded|muttered|mutters|exclaimed|exclaims|added|adds|'
    r'followed|follows|stuttered|stutters|mentioned|mentions|'
    r'questioned|questions|stated|states|declared|declares)\b'
)

_POSSESSIVE = re.compile(r'^([A-Z][a-z]{2,})\'s\b', re.MULTILINE)


def find_characters_in_script(script_text: str) -> list:
    """
    Scan script_text for character names and return a sorted unique list.
    Always includes 'Narrator'.

    Three passes:
      1. Explicit [Name] tags
      2. Prose attribution — "Name said / replied / asked / …"
      3. Possessives at line start — "Name's hands blurred…"
    """
    names = set()
    names.add("Narrator")

    # Pass 1: tagged lines [Name] or [Name]:
    for m in re.finditer(r'^\[([^\]]+)\]', script_text, re.MULTILINE):
        name = m.group(1).strip()
        if name != "Narrator":
            names.add(name)

    # Pass 2: "Name said / replied / …" attribution
    for m in _ATTR_VERBS.finditer(script_text):
        candidate = m.group(1)
        if candidate not in _SKIP_WORDS:
            names.add(candidate)

    # Pass 3: "Name's …" possessives at the start of a line
    for m in _POSSESSIVE.finditer(script_text):
        candidate = m.group(1)
        if candidate not in _SKIP_WORDS:
            names.add(candidate)

    return sorted(names)


def tag_script_with_ai(
    source_text: str,
    engine: str,
    characters: list,
    on_progress=None,
) -> str:
    """
    Use the local LLM to convert prose to a [Character] tagged script.

    engine     : 'kokoro' | 'fish14' | 'fish15' | 's1mini' | 's1'
    characters : list of known character name strings
    on_progress: optional callable(message: str, fraction: float)

    Returns the tagged script as a string.
    """
    try:
        from tag_suggester import (
            _infer_chunk, is_llm_available, is_qwen_model_ready,
            _load_llm, _active_is_ollama, _llm_lock,
        )
    except ImportError as exc:
        raise RuntimeError(f"LLM not available: {exc}")

    if not is_llm_available():
        raise RuntimeError(
            "LLM runtime not installed.\n"
            "Install llama-cpp-python from Settings."
        )
    if not is_qwen_model_ready():
        raise RuntimeError(
            "No LLM model downloaded.\n"
            "Download a model from Settings."
        )

    char_list = ", ".join(characters) if characters else "unknown"
    # Use editable prompt from tag_suggester if available, else fall back to built-in
    try:
        from tag_suggester import get_prompt as _get_prompt
        _key = "script_kokoro" if engine == "kokoro" else "script_fish"
        _base = _get_prompt(_key) or (_KOKORO_PROMPT if engine == "kokoro" else _FISH_PROMPT)
    except Exception:
        _base = _KOKORO_PROMPT if engine == "kokoro" else _FISH_PROMPT
    system_prompt = _base.format(characters=char_list)

    if not _active_is_ollama():
        with _llm_lock:
            _load_llm()

    paragraphs = [p.strip() for p in re.split(r'\n{2,}', source_text) if p.strip()]
    total = max(len(paragraphs), 1)
    result_parts = []

    def _split_para(para: str, limit: int = 1500) -> list:
        """Split a long paragraph at sentence boundaries, not arbitrary char positions."""
        if len(para) <= limit:
            return [para]
        # Split on sentence-ending punctuation followed by whitespace
        sentences = re.split(r'(?<=[.!?"])\s+', para)
        chunks, current = [], ""
        for sent in sentences:
            if current and len(current) + len(sent) + 1 > limit:
                chunks.append(current.strip())
                current = sent
            else:
                current = (current + " " + sent).strip() if current else sent
        if current:
            chunks.append(current.strip())
        return chunks or [para]

    for i, para in enumerate(paragraphs):
        if on_progress:
            on_progress(f"Tagging paragraph {i + 1} of {total}...", i / total)

        chunks = _split_para(para)
        for chunk in chunks:
            try:
                max_tok = min(1200, int(len(chunk) / 4 * 3) + 300)
                tagged = _infer_chunk(
                    system_prompt, chunk,
                    max_tokens=max_tok,
                    temperature=0.2,
                    top_p=0.95,
                    repeat_penalty=1.1,
                )
                if tagged and tagged.strip():
                    result_parts.append(tagged.strip())
                else:
                    result_parts.append(f"[Narrator] {chunk}")
            except Exception as exc:
                logger.warning("tag_script chunk failed: %s", exc)
                result_parts.append(f"[Narrator] {chunk}")

    if on_progress:
        on_progress("Done.", 1.0)

    return "\n".join(result_parts)


def enhance_script_flow(
    script_text: str,
    engine: str,
    on_progress=None,
) -> str:
    """
    Run the already-tagged script through the LLM to improve emotional
    continuity, natural delivery, and conversation flow.

    engine     : 'kokoro' | 'fish14' | 'fish15' | 's1mini' | 's1'
    on_progress: optional callable(message: str, fraction: float)

    Returns the enhanced script as a string.
    """
    try:
        from tag_suggester import (
            _infer_chunk, is_llm_available, is_qwen_model_ready,
            _load_llm, _active_is_ollama, _llm_lock,
        )
    except ImportError as exc:
        raise RuntimeError(f"LLM not available: {exc}")

    if not is_llm_available():
        raise RuntimeError(
            "LLM runtime not installed.\n"
            "Install llama-cpp-python from Settings."
        )
    if not is_qwen_model_ready():
        raise RuntimeError(
            "No LLM model downloaded.\n"
            "Download a model from Settings."
        )

    system_prompt = _ENHANCE_KOKORO_PROMPT if engine == "kokoro" else _ENHANCE_FISH_PROMPT

    if not _active_is_ollama():
        with _llm_lock:
            _load_llm()

    # Split into ~30-line windows so the model sees enough context for continuity
    lines = [l for l in script_text.splitlines() if l.strip()]
    window = 30
    result_lines = []
    total = max(len(lines), 1)

    for i in range(0, len(lines), window):
        chunk_lines = lines[i:i + window]
        chunk = "\n".join(chunk_lines)
        if on_progress:
            on_progress(
                f"Enhancing lines {i + 1}–{min(i + window, total)} of {total}...",
                i / total,
            )
        try:
            max_tok = min(1200, int(len(chunk) / 4 * 3) + 300)
            enhanced = _infer_chunk(
                system_prompt, chunk,
                max_tokens=max_tok,
                temperature=0.3,
                top_p=0.95,
                repeat_penalty=1.1,
            )
            if enhanced and enhanced.strip():
                result_lines.extend(enhanced.strip().splitlines())
            else:
                result_lines.extend(chunk_lines)
        except Exception as exc:
            logger.warning("enhance_script_flow chunk failed: %s", exc)
            result_lines.extend(chunk_lines)

    if on_progress:
        on_progress("Enhancement complete.", 1.0)

    return "\n".join(result_lines)
